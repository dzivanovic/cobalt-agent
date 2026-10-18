"""F23 playbook renderer, offline (card 503 rows P1-P5; FINAL §5 as amended by X6).

THE FIXTURE IS SYNTHETIC (L32, L45). The vault is a `tmp_path` folder of
strategy notes built from the one note the repo ships
(`configs/cobalt/taxonomy/examples/example_trade_def.md`), re-keyed to the
F-15 walk's slugs: `od-only` `[opening_drive]`, `both` `[opening_drive,
continuation]`, and the real draft `draft-x` (an empty Definition unit — a
draft the loader lists, never a finished `family: []`). The radar rows are
the walk's ids 1-14, constructed tickers only. The store and `job_run` are
faked at the module seam: nothing here reaches `cobalt_dev`.

Every `cobalt.playbook` import sits inside a test, so on a tree without the
package each test is red for its own reason, not a collection error.
"""

from __future__ import annotations

import ast
import os
import subprocess
import sys
import tomllib
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest
from pydantic import ValidationError

from cobalt.taxonomy.loader import EXAMPLE_NOTE_PATH
from cobalt.taxonomy.slug import trade_key
from cobalt.taxonomy.vault_loader import STRATEGIES_DIR

REPO = Path(__file__).resolve().parents[2]
ET = ZoneInfo("America/New_York")
RUN_DATE = "2026-10-11"  # the walk's Sunday
WALK_FORMED = datetime(2026, 10, 5, 0, 1, tzinfo=ET)
SIX = ("opening_drive", "continuation", "range_break", "reversion", "time_window", "unfiled")

DRAFT_NOTE = """---
trade_def: draft-x
name: Draft X
status: draft
---
## Definition
<!-- cobalt:section definition -->
<!-- cobalt:unit trade_def:draft-x -->
<!-- /cobalt:unit trade_def:draft-x -->
<!-- /cobalt:section definition -->
"""


# ---------------------------------------------------------------------
# the synthetic vault and rows (shared with test_playbook_db.py)
# ---------------------------------------------------------------------


def note_text(slug: str, name: str, families: list[str]) -> str:
    """The shipped example note, re-keyed to `slug` with `families` — the
    same rewrite a trader does by hand when copying a strategy."""
    text = EXAMPLE_NOTE_PATH.read_text(encoding="utf-8")
    return (
        text.replace("example-range-break", slug)
        .replace("example_range_break", trade_key(slug))
        .replace("Example Range Break", name)
        .replace("family: [range_break]", "family: [" + ", ".join(families) + "]")
    )


def make_vault(root: Path, notes: dict[str, str]) -> Path:
    strategies = root / STRATEGIES_DIR
    strategies.mkdir(parents=True, exist_ok=True)
    for name, text in notes.items():
        (strategies / f"{name}.md").write_text(text, encoding="utf-8")
    return root


def walk_vault(root: Path) -> Path:
    return make_vault(root, {
        "OD Only": note_text("od-only", "OD Only", ["opening_drive"]),
        "Both": note_text("both", "Both Ways", ["opening_drive", "continuation"]),
        "Draft X": DRAFT_NOTE,
    })


def radar_row(card_id: int, slug: str, *, formed_at=WALK_FORMED, state: str = "WATCH",
              why: str | None = "synthetic why", card_score: int | None = None) -> dict:
    """One row as the playbook's SELECT returns it (synthetic, L32)."""
    return {
        "id": card_id, "ticker": f"ZZP{chr(64 + card_id % 26 or 26)}", "direction": "long", "state": state,
        "trade_def_slug": slug, "formed_at": formed_at, "sized_grade": None, "tapped_grade": None,
        "trigger_price": Decimal("5.5000"), "structural_stop": Decimal("5.2000"), "entry": Decimal("5.5000"),
        "stop": Decimal("5.2000"), "card_score": card_score, "why": why,
    }


def walk_rows() -> list[dict]:
    """F-15: ids 1-12 `od-only` (id 12 EXPIRED, id 1 a 300-char why), id 13
    `both` FILLED a day after the titled week, id 14 `gone` (no note)."""
    rows = [radar_row(i, "od-only") for i in range(1, 13)]
    rows[0]["why"] = "w" * 300
    rows[11]["state"] = "EXPIRED"
    rows.append(radar_row(13, "both", state="FILLED", formed_at=datetime(2026, 10, 12, 0, 1, tzinfo=ET)))
    rows.append(radar_row(14, "gone"))
    return rows


def _book(vault: Path, rows: list[dict], latest_per_family=None):
    from cobalt.playbook import data
    from cobalt.taxonomy.vault_loader import load_vault_trade_defs

    return data.group(load_vault_trade_defs(vault), rows, latest_per_family=latest_per_family)


def _write(book, folder: Path, *, rows_per_slide: int = 8, why_max_chars: int = 120):
    from cobalt.playbook import deck

    folder.mkdir(parents=True, exist_ok=True)
    return deck.write_all(
        book, folder, run_date=date(2026, 10, 11), rendered_at=datetime(2026, 10, 11, 17, 30, tzinfo=ET),
        rows_per_slide=rows_per_slide, why_max_chars=why_max_chars,
    )


def _slides(path: Path) -> list:
    from pptx import Presentation

    return list(Presentation(str(path)).slides)


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if shape.has_table:
            parts.extend(" | ".join(c.text for c in r.cells) for r in shape.table.rows)
    return "\n".join(parts)


def _table_slides(path: Path) -> list[list[list[str]]]:
    """Per table slide, its data rows (header dropped), each a list of cell texts."""
    out = []
    for slide in _slides(path):
        for shape in slide.shapes:
            if shape.has_table:
                out.append([[c.text for c in r.cells] for r in list(shape.table.rows)[1:]])
    return out


@pytest.fixture
def vault(tmp_path):
    return walk_vault(tmp_path / "vault")


# ---------------------------------------------------------------------
# P1 — the pin
# ---------------------------------------------------------------------


def test_python_pptx_is_pinned_exact():
    deps = tomllib.loads((REPO / "pyproject.toml").read_text())["project"]["dependencies"]
    named = [d for d in deps if d.replace(" ", "").lower().startswith("python-pptx")]
    assert len(named) == 1 and named[0].startswith("python-pptx=="), (
        f"pyproject.toml must hold exactly one exact pin `python-pptx==<version>`, got {named}"
    )


# ---------------------------------------------------------------------
# P2 — the config
# ---------------------------------------------------------------------


def test_config_loads_the_committed_file():
    from cobalt.playbook.config import load_playbook_config

    cfg = load_playbook_config()
    assert (cfg.latest_per_family, cfg.rows_per_slide, cfg.why_max_chars, cfg.out_dir) == (
        None, 8, 120, "data/playbook",
    )


GOOD = {"latest_per_family": None, "rows_per_slide": 8, "why_max_chars": 120, "out_dir": "data/playbook"}


@pytest.mark.parametrize("raw", [
    pytest.param({**GOOD, "font": "Arial"}, id="unknown-key"),
    pytest.param({**GOOD, "rows_per_slide": "8"}, id="rows-per-slide-a-string"),
    pytest.param({**GOOD, "why_max_chars": 120.5}, id="why-max-chars-a-float"),
    pytest.param({**GOOD, "latest_per_family": "all"}, id="latest-per-family-a-word"),
    pytest.param({**GOOD, "out_dir": 5}, id="out-dir-a-number"),
    pytest.param({k: v for k, v in GOOD.items() if k != "rows_per_slide"}, id="missing-key-no-default"),
])
def test_config_refuses_an_unknown_key_and_a_bad_type(tmp_path, raw):
    import yaml

    from cobalt.playbook.config import load_playbook_config

    path = tmp_path / "playbook.yaml"
    path.write_text(yaml.safe_dump(raw))
    with pytest.raises(ValidationError):
        load_playbook_config(path)


# ---------------------------------------------------------------------
# P3 — the data
# ---------------------------------------------------------------------


def _ids(deck) -> list[int]:
    return [c.id for c in deck.cards]


def test_two_family_card_is_on_both_decks_with_k_1(vault):
    book = _book(vault, walk_rows())
    od, cont = book.deck("opening_drive"), book.deck("continuation")
    assert 13 in _ids(od) and 13 in _ids(cont)
    assert od.also_elsewhere == 1 and cont.also_elsewhere == 1
    assert [d.name for d in book.decks] == list(SIX)


def test_a_draft_is_listed_and_in_no_deck(vault):
    book = _book(vault, [*walk_rows(), radar_row(15, "draft-x")])
    assert book.drafts == ("draft-x",)
    for deck in book.decks:
        assert "draft-x" not in [s.slug for s in deck.setups]
        if deck.name != "unfiled":
            assert 15 not in _ids(deck)
    assert {c.id: c.reason for c in book.deck("unfiled").cards}[15] == "draft"


def test_no_cap_and_null_formed_at_sorts_last_and_is_present(vault):
    rows = [radar_row(i, "od-only", formed_at=WALK_FORMED + timedelta(minutes=i)) for i in range(1, 13)]
    rows.append(radar_row(99, "od-only", formed_at=None))  # the highest id, no time: last, never dropped
    book = _book(vault, rows, latest_per_family=None)
    assert _ids(book.deck("opening_drive")) == [*range(12, 0, -1), 99]
    assert book.deck("opening_drive").in_family == 13


def test_unfiled_rows_carry_slug_and_reason(vault):
    rows = [radar_row(14, "gone"), radar_row(15, "draft-x", formed_at=WALK_FORMED + timedelta(hours=1))]
    unfiled = _book(vault, rows).deck("unfiled")
    assert [(c.id, c.slug, c.reason) for c in unfiled.cards] == [(15, "draft-x", "draft"), (14, "gone", "no note")]
    assert unfiled.in_family == 2


def test_card_score_never_orders(vault):
    """Same `formed_at`, scores reversed both ways: `id DESC` decides (L52)."""
    for low_first in (True, False):
        rows = [radar_row(1, "od-only", card_score=90 if low_first else 10),
                radar_row(2, "od-only", card_score=10 if low_first else 90),
                radar_row(3, "od-only", card_score=None)]
        assert _ids(_book(vault, rows).deck("opening_drive")) == [3, 2, 1]


# ---------------------------------------------------------------------
# P4 — the deck
# ---------------------------------------------------------------------


def test_walk_slide_counts(vault, tmp_path):
    written = _write(_book(vault, walk_rows()), tmp_path / "out")
    counts = {w.name: len(_slides(w.path)) for w in written}
    assert counts == {"opening_drive": 5, "continuation": 3, "range_break": 2, "reversion": 2,
                      "time_window": 2, "unfiled": 2}
    assert {w.name: w.slides for w in written} == counts
    assert sorted(p.name for p in (tmp_path / "out").iterdir()) == sorted(f"{n}.pptx" for n in SIX)


def test_table_rows_newest_first_8_per_slide(vault, tmp_path):
    rows = [radar_row(i, "od-only", formed_at=WALK_FORMED + timedelta(minutes=i)) for i in range(1, 18)]
    written = {w.name: w for w in _write(_book(vault, rows), tmp_path / "out")}
    tables = _table_slides(written["opening_drive"].path)
    assert [len(t) for t in tables] == [8, 8, 1]
    tickers = [r[1] for t in tables for r in t]
    assert tickers == [radar_row(i, "od-only")["ticker"] for i in range(17, 0, -1)]
    dates = [r[0] for t in tables for r in t]
    assert dates == sorted(dates, reverse=True)


def test_empty_family_deck_says_so(vault, tmp_path):
    written = {w.name: w for w in _write(_book(vault, [radar_row(1, "od-only")]), tmp_path / "out")}
    assert "no setup in this family" in _slide_text(_slides(written["time_window"].path)[1])
    assert "no cards yet" in _slide_text(_slides(written["continuation"].path)[-1])
    assert "no unfiled cards" in _slide_text(_slides(written["unfiled"].path)[1])


def test_why_is_cut_to_120_plus_ellipsis(vault, tmp_path):
    written = {w.name: w for w in _write(_book(vault, walk_rows()), tmp_path / "out")}
    rows = [r for t in _table_slides(written["opening_drive"].path) for r in t]
    by_ticker = {r[1]: r for r in rows}
    why = by_ticker[radar_row(1, "od-only")["ticker"]][11]
    assert why == "w" * 120 + "…" and len(why) == 121
    assert by_ticker[radar_row(2, "od-only")["ticker"]][11] == "synthetic why"


def test_both_titles_say_also_in_another_family_1(vault, tmp_path):
    written = {w.name: w for w in _write(_book(vault, walk_rows()), tmp_path / "out")}
    for name in ("opening_drive", "continuation"):
        title = _slide_text(_slides(written[name].path)[0])
        assert "also in another family: 1" in title
        assert "cards are the newest by formed_at, not limited to this week" in title
        assert "Mon–Fri ET ending this Sunday" in title and "2026-10-05" in title and "2026-10-09" in title


# ---------------------------------------------------------------------
# P5 — the command
# ---------------------------------------------------------------------


class FakeJobs:
    """`job_run` faked at the module seam: one (label, exit_code, result) per run."""

    def __init__(self):
        self.rows: list[tuple] = []

    def __call__(self, label, **_kw):
        @contextmanager
        def run():
            handle = SimpleNamespace(label=label, result=None)
            try:
                yield handle
            except BaseException:
                self.rows.append((label, 1, handle.result))
                raise
            self.rows.append((label, 0, handle.result))

        return run()


@pytest.fixture
def seam(monkeypatch, vault):
    """The store and `job_run` faked; the vault is the walk's."""
    from cobalt.playbook import cli as playbook_cli
    from cobalt.playbook import data

    monkeypatch.setenv("COBALT_VAULT_PATH", str(vault))
    store = SimpleNamespace(rows=walk_rows(), dbnames=[])

    def fetch(dbname):
        store.dbnames.append(dbname)
        return [dict(r) for r in store.rows]

    jobs = FakeJobs()
    monkeypatch.setattr(data, "fetch_radar_rows", fetch)
    monkeypatch.setattr(playbook_cli, "job_run", jobs)
    return SimpleNamespace(store=store, jobs=jobs, vault=vault)


def _cli(monkeypatch, *argv) -> int:
    import cobalt.cli

    monkeypatch.setattr(sys, "argv", ["cobalt", *argv])
    try:
        cobalt.cli.main()
    except SystemExit as e:
        return e.code
    return 0


def _render(monkeypatch, out: Path, *extra) -> int:
    return _cli(monkeypatch, "playbook", "render", "--date", RUN_DATE, "--out", str(out), *extra)


def _snapshot(folder: Path) -> dict[str, tuple[int, bytes]]:
    return {str(p.relative_to(folder)): (p.stat().st_mtime_ns, p.read_bytes())
            for p in sorted(folder.rglob("*")) if p.is_file()}


def test_rerun_same_date_leaves_only_the_new_run(seam, monkeypatch, tmp_path, capsys):
    out = tmp_path / "out"
    assert _render(monkeypatch, out) == 0
    first = capsys.readouterr().out
    for name in SIX:
        assert str(out / RUN_DATE / f"{name}.pptx") in first
    assert "unfiled cards: 1" in first and "drafts: draft-x" in first
    assert seam.jobs.rows[-1][:2] == ("com.cobalt.playbook", 0)
    assert seam.jobs.rows[-1][2]["families"]["opening_drive"]["cards"] == 13

    (out / f"{RUN_DATE}.tmp").mkdir()
    (out / f"{RUN_DATE}.tmp" / "half-written.pptx").write_bytes(b"leftover")
    seam.store.rows = [r for r in walk_rows() if r["id"] != 14]  # the new run has no unfiled card
    assert _render(monkeypatch, out) == 0
    assert sorted(p.name for p in out.iterdir()) == [RUN_DATE]
    assert sorted(p.name for p in (out / RUN_DATE).iterdir()) == sorted(f"{n}.pptx" for n in SIX)
    assert "no unfiled cards" in _slide_text(_slides(out / RUN_DATE / "unfiled.pptx")[1])


def test_a_crash_before_the_swap_keeps_the_previous_folder(seam, monkeypatch, tmp_path):
    from cobalt.playbook import cli as playbook_cli

    out = tmp_path / "out"
    assert _render(monkeypatch, out) == 0
    before = _snapshot(out / RUN_DATE)
    real = playbook_cli._rename

    def crash_on_the_second(src: Path, dst: Path) -> None:
        if src.name == f"{RUN_DATE}.tmp":
            raise OSError("synthetic crash between the two renames")
        real(src, dst)

    monkeypatch.setattr(playbook_cli, "_rename", crash_on_the_second)
    assert _render(monkeypatch, out) == 1
    assert _snapshot(out / RUN_DATE) == before
    assert not (out / f"{RUN_DATE}.old").exists()
    assert seam.jobs.rows[-1][:2] == ("com.cobalt.playbook", 1)


def test_validation_error_exits_1_and_keeps_the_previous_folder(seam, monkeypatch, tmp_path, capsys):
    out = tmp_path / "out"
    assert _render(monkeypatch, out) == 0
    before = _snapshot(out / RUN_DATE)
    note = seam.vault / STRATEGIES_DIR / "OD Only.md"
    note.write_text(note.read_text().replace("entry_mode: front_side", "entry_mode: sideways"))
    capsys.readouterr()
    assert _render(monkeypatch, out) == 1
    assert "VaultTaxonomyError" in capsys.readouterr().err
    assert seam.jobs.rows[-1][:2] == ("com.cobalt.playbook", 1)
    assert _snapshot(out / RUN_DATE) == before
    assert sorted(p.name for p in out.iterdir()) == [RUN_DATE]


def test_dry_run_writes_no_deck_and_no_row(seam, monkeypatch, tmp_path, capsys):
    out = tmp_path / "out"
    assert _render(monkeypatch, out, "--dry-run") == 0
    printed = capsys.readouterr().out
    assert not out.exists()
    assert seam.jobs.rows == []
    for name in SIX:
        assert name in printed and str(out / RUN_DATE / f"{name}.pptx") in printed
    assert "DRY RUN" in printed


def test_no_vault_write(seam, monkeypatch, tmp_path):
    before = _snapshot(seam.vault)
    assert _render(monkeypatch, tmp_path / "out") == 0
    assert _snapshot(seam.vault) == before
    reached = []
    for path in sorted((REPO / "src" / "cobalt" / "playbook").glob("*.py")):
        for node in ast.walk(ast.parse(path.read_text())):
            if isinstance(node, ast.Import):
                reached += [(path.name, a.name) for a in node.names if a.name.startswith("cobalt.vaultwrite")]
            elif isinstance(node, ast.ImportFrom) and (node.module or "").startswith("cobalt.vaultwrite"):
                reached.append((path.name, node.module))
    assert reached == [], f"cobalt.playbook imports the vault writer: {reached}"


def test_import_cobalt_cli_leaves_pptx_unloaded():
    probe = "import sys, cobalt.cli; print('pptx' in sys.modules, 'cobalt.playbook.cli' in sys.modules)"
    done = subprocess.run([sys.executable, "-c", probe], capture_output=True, text=True, cwd=REPO,
                          env={**os.environ, "COBALT_ENV": "dev"})
    assert done.returncode == 0, done.stderr
    assert done.stdout.split() == ["False", "True"]


def test_the_cli_mounts_playbook_render(monkeypatch, capsys):
    code = _cli(monkeypatch, "playbook", "render", "--help")
    printed = capsys.readouterr()
    assert code == 0, printed.err
    for flag in ("--date", "--out", "--dry-run"):
        assert flag in printed.out
