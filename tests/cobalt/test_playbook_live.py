"""F23 card 503 row P8: the render against his strategies folder (live-note suite).

Reads `COBALT_LIVE_VAULT_ROOT` (read only) and writes six decks into
`tmp_path`; no deck is committed and every count of his stays in this
process (L32). The live-note pass runs with `.env` absent (BUILD-HUB W (e)),
so `cobalt_dev` is unreachable there: the card store is faked empty at the
module seam, and the cards side is `test_playbook_db.py`'s. What this proves
is the vault side on real notes: six files, and each family deck's setup
slides equal the defs the loader returns for that family.
"""

from __future__ import annotations

import os
from pathlib import Path

import pytest

LIVE = os.getenv("COBALT_LIVE_VAULT_ROOT")
requires_live = pytest.mark.skipif(
    not LIVE, reason="COBALT_LIVE_VAULT_ROOT not set — the hub runs the live playbook render"
)

SIX = ("opening_drive", "continuation", "range_break", "reversion", "time_window", "unfiled")


def _snapshot(folder: Path) -> dict[str, int]:
    return {str(p): p.stat().st_mtime_ns for p in sorted(folder.rglob("*.md"))}


@requires_live
def test_live_vault_renders_six_decks(monkeypatch, tmp_path):
    from cobalt.playbook import cli as playbook_cli
    from cobalt.playbook import data
    from cobalt.taxonomy.vault_loader import STRATEGIES_DIR, load_vault_trade_defs
    from pptx import Presentation

    root = Path(LIVE)
    before = _snapshot(root / STRATEGIES_DIR)
    monkeypatch.setenv("COBALT_VAULT_PATH", str(root))
    monkeypatch.setenv("COBALT_ALLOW_DEV_ENTRY", "1")  # a read of the live vault from a dev run, on purpose
    monkeypatch.setattr(data, "fetch_radar_rows", lambda dbname: [])
    runs = []

    class _Run:
        result = None

    from contextlib import contextmanager

    @contextmanager
    def fake_job_run(label, **_kw):
        handle = _Run()
        yield handle
        runs.append((label, handle.result))

    monkeypatch.setattr(playbook_cli, "job_run", fake_job_run)
    out = tmp_path / "out"
    args = playbook_cli.build_parser().parse_args(["render", "--date", "2026-10-11", "--out", str(out)])
    args.func(args)

    folder = out / "2026-10-11"
    assert sorted(p.name for p in folder.iterdir()) == sorted(f"{n}.pptx" for n in SIX)
    defs = load_vault_trade_defs(root).defs
    for family in SIX[:-1]:
        expected = sum(1 for d in defs if family in {f.value for f in d.definition.family})
        slides = Presentation(str(folder / f"{family}.pptx")).slides
        setup_slides = [s for s in slides if any(
            sh.has_text_frame and sh.text_frame.text.startswith("setup · ") for sh in s.shapes)]
        assert len(setup_slides) == expected, family
    assert _snapshot(root / STRATEGIES_DIR) == before
    assert [label for label, _ in runs] == ["com.cobalt.playbook"]
