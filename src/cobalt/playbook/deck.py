"""F23 §2.2: the decks — `data.Deck` in, a `.pptx` file out.

A plain python-pptx presentation built in code (FINAL C4): blank 16:9,
default theme, no template, no image, no chart. The only I/O is the save.

  1. Title slide: the family, the week label (`Mon–Fri ET ending this
     Sunday`), the subtitle that says the cards are not limited to that
     week, `rendered HH:MM ET`, the counts and `also in another family: K`.
  2. One slide per setup: name, note path, class, its families, its cards.
  3. Table slides, `rows_per_slide` cards each, newest first, the twelve
     columns of §2.2.3; `why` cut to `why_max_chars` plus `…`.
  4. An empty deck still has a content slide: `no setup in this family`,
     `no cards yet` (setups but no card), or `no unfiled cards`.

`card_score` is printed as stored; nothing here orders anything (L52) —
the cards arrive in `data.newest_first` order and keep it.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Optional

from pptx import Presentation
from pptx.util import Inches, Pt

from .data import Card, Deck, Playbook

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
#: The default template's "Blank" layout.
BLANK_LAYOUT = 6
MARGIN = Inches(0.4)
#: FINAL §2.2.3, in order.
COLUMNS = (
    "date (ET)", "ticker", "setup", "direction", "state", "grade", "trigger", "structural stop",
    "entry", "stop", "card_score", "why",
)
#: Relative column widths; `why` takes what is left.
_WIDTHS = (1.25, 0.7, 1.4, 0.7, 0.85, 0.55, 0.75, 0.8, 0.7, 0.7, 0.7, 4.53)
NO_SETUP = "no setup in this family"
NO_CARDS = "no cards yet"
NO_UNFILED = "no unfiled cards"
SUBTITLE = "cards are the newest by formed_at, not limited to this week"
SETUP_PREFIX = "setup · "
DASH = "—"


@dataclass(frozen=True)
class Written:
    name: str
    path: Path
    slides: int
    cards: int


def week_label(run_date: date) -> str:
    """Mon–Fri of the week whose Sunday is `run_date` (or the next Sunday)."""
    sunday = run_date + timedelta(days=6 - run_date.weekday())
    return f"Mon–Fri ET ending this Sunday: {sunday - timedelta(days=6)} – {sunday - timedelta(days=2)}"


def cut(text: Optional[str], limit: int) -> str:
    if not text:
        return DASH
    return text if len(text) <= limit else text[:limit] + "…"


def _value(value: Any) -> str:
    return DASH if value is None else str(value)


def _lines(slide, lines: list[str], *, top, height, size: int, first_size: Optional[int] = None) -> None:
    box = slide.shapes.add_textbox(MARGIN, top, SLIDE_WIDTH - 2 * MARGIN, height)
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(first_size if (i == 0 and first_size) else size)


def _title_slide(prs, deck: Deck, *, run_date: date, rendered_at: datetime) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    if deck.is_unfiled:
        counts = [f"unfiled cards {len(deck.cards)} — each row carries its slug and the reason (draft / no note)"]
    else:
        counts = [
            f"setups {len(deck.setups)} · cards shown {len(deck.cards)} · cards in family {deck.in_family}",
            f"also in another family: {deck.also_elsewhere}",
        ]
    _lines(slide, [deck.name, week_label(run_date), SUBTITLE, f"rendered {rendered_at:%H:%M} ET", *counts],
           top=Inches(1.5), height=Inches(4.5), size=20, first_size=40)


def _setup_slide(prs, setup) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    _lines(slide, [
        f"{SETUP_PREFIX}{setup.name}", setup.note_path, f"class {setup.trade_class}",
        f"families {', '.join(setup.families)}", f"cards in this deck {setup.cards}",
    ], top=Inches(1.5), height=Inches(4.0), size=20, first_size=32)


def _message_slide(prs, text: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    _lines(slide, [text], top=Inches(3.0), height=Inches(1.5), size=28)


def _row(card: Card, *, to_et, why_max_chars: int) -> list[str]:
    setup = card.setup_name if card.reason is None else f"{card.slug} · {card.reason}"
    when = f"{to_et(card.formed_at):%Y-%m-%d %H:%M}" if card.formed_at is not None else DASH
    return [
        when, card.ticker, setup, card.direction, card.state, _value(card.grade), _value(card.trigger_price),
        _value(card.structural_stop), _value(card.entry), _value(card.stop), _value(card.card_score),
        cut(card.why, why_max_chars),
    ]


def _table_slide(prs, deck: Deck, chunk: list[Card], first: int, *, to_et, why_max_chars: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    _lines(slide, [f"{deck.name} · cards {first + 1}–{first + len(chunk)} of {len(deck.cards)}"],
           top=Inches(0.2), height=Inches(0.6), size=18)
    width = SLIDE_WIDTH - 2 * MARGIN
    table = slide.shapes.add_table(len(chunk) + 1, len(COLUMNS), MARGIN, Inches(0.9), width,
                                   Inches(0.4) * (len(chunk) + 1)).table
    total = sum(_WIDTHS)
    for i, share in enumerate(_WIDTHS):
        table.columns[i].width = int(width * share / total)
    for r, cells in enumerate([list(COLUMNS)] + [_row(c, to_et=to_et, why_max_chars=why_max_chars) for c in chunk]):
        for c, text in enumerate(cells):
            cell = table.cell(r, c)
            cell.text = text
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(9)


def build(deck: Deck, *, run_date: date, rendered_at: datetime, rows_per_slide: int, why_max_chars: int):
    """The presentation for one deck (not saved)."""
    from cobalt.session.clock import session_clock

    to_et = session_clock().to_et
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_WIDTH, SLIDE_HEIGHT
    _title_slide(prs, deck, run_date=run_date, rendered_at=rendered_at)
    for setup in deck.setups:
        _setup_slide(prs, setup)
    if deck.is_unfiled and not deck.cards:
        _message_slide(prs, NO_UNFILED)
    elif not deck.is_unfiled and not deck.setups:
        _message_slide(prs, NO_SETUP)
    elif not deck.cards:
        _message_slide(prs, NO_CARDS)
    cards = list(deck.cards)
    for first in range(0, len(cards), rows_per_slide):
        _table_slide(prs, deck, cards[first:first + rows_per_slide], first, to_et=to_et,
                     why_max_chars=why_max_chars)
    return prs


def write_deck(deck: Deck, path: Path, *, run_date: date, rendered_at: datetime, rows_per_slide: int,
               why_max_chars: int) -> int:
    """Save one deck to `path`; returns its slide count. A save error raises (L1)."""
    prs = build(deck, run_date=run_date, rendered_at=rendered_at, rows_per_slide=rows_per_slide,
                why_max_chars=why_max_chars)
    prs.save(str(path))
    return len(prs.slides)


def write_all(book: Playbook, folder: Path, *, run_date: date, rendered_at: datetime, rows_per_slide: int,
              why_max_chars: int) -> list[Written]:
    """All six decks into `folder`, `<deck>.pptx` each."""
    written = []
    for deck in book.decks:
        path = folder / f"{deck.name}.pptx"
        slides = write_deck(deck, path, run_date=run_date, rendered_at=rendered_at,
                            rows_per_slide=rows_per_slide, why_max_chars=why_max_chars)
        written.append(Written(name=deck.name, path=path, slides=slides, cards=len(deck.cards)))
    return written


__all__ = ["COLUMNS", "NO_CARDS", "NO_SETUP", "NO_UNFILED", "SUBTITLE", "Written", "build", "cut",
           "week_label", "write_all", "write_deck"]
